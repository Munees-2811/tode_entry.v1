import json
import csv
import os
from datetime import datetime


class DataManager:
    def __init__(self):
        self.entries = []
        self.field_names = []
        self.current_entry = {}
        self.templates = {}
        self._load_default_templates()

    def _load_default_templates(self):
        self.templates = {
            "Invoice": ["Invoice No", "Date", "Vendor", "Amount", "Tax", "Total", "Notes"],
            "Receipt": ["Store", "Date", "Items", "Subtotal", "Tax", "Total", "Payment Method"],
            "Form": ["Name", "Date", "Address", "Phone", "Email", "ID Number", "Notes"],
            "Medical": ["Patient Name", "Date", "Doctor", "Diagnosis", "Prescription", "Notes"],
            "Custom": ["Field 1", "Field 2", "Field 3", "Field 4", "Field 5"],
        }

    def get_template_names(self):
        return list(self.templates.keys())

    def get_template_fields(self, name):
        return self.templates.get(name, [])

    def save_template(self, name, fields):
        self.templates[name] = list(fields)

    def set_field_names(self, names):
        self.field_names = list(names)

    def update_field(self, field_name, value):
        self.current_entry[field_name] = value

    def save_current_entry(self):
        if self.current_entry:
            entry = dict(self.current_entry)
            entry["_timestamp"] = datetime.now().isoformat()
            self.entries.append(entry)
            self.current_entry = {}
            return True
        return False

    def delete_entry(self, index):
        if 0 <= index < len(self.entries):
            self.entries.pop(index)
            return True
        return False

    def edit_entry(self, index, field, value):
        if 0 <= index < len(self.entries):
            self.entries[index][field] = value
            return True
        return False

    def get_all_field_names(self):
        all_keys = []
        for entry in self.entries:
            for k in entry:
                if k not in all_keys and not k.startswith("_"):
                    all_keys.append(k)
        return all_keys

    def export_csv(self, filepath):
        if not self.entries:
            return False

        all_keys = []
        for entry in self.entries:
            for k in entry:
                if k not in all_keys:
                    all_keys.append(k)

        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=all_keys)
            writer.writeheader()
            writer.writerows(self.entries)
        return True

    def export_json(self, filepath):
        if not self.entries:
            return False

        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(self.entries, f, indent=2, ensure_ascii=False)
        return True

    def export_excel(self, filepath):
        if not self.entries:
            return False

        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl not installed. Run: pip install openpyxl")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TODE Data"

        all_keys = []
        for entry in self.entries:
            for k in entry:
                if k not in all_keys:
                    all_keys.append(k)

        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        header_fill = openpyxl.styles.PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        header_align = openpyxl.styles.Alignment(horizontal="center")
        thin_border = openpyxl.styles.Border(
            left=openpyxl.styles.Side(style="thin"),
            right=openpyxl.styles.Side(style="thin"),
            top=openpyxl.styles.Side(style="thin"),
            bottom=openpyxl.styles.Side(style="thin"),
        )

        for col_idx, key in enumerate(all_keys, 1):
            cell = ws.cell(row=1, column=col_idx, value=key)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        alt_fill = openpyxl.styles.PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
        for row_idx, entry in enumerate(self.entries, 2):
            for col_idx, key in enumerate(all_keys, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=entry.get(key, ""))
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

        for col_idx, key in enumerate(all_keys, 1):
            max_len = len(key)
            for entry in self.entries:
                val = str(entry.get(key, ""))
                max_len = max(max_len, len(val))
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 50)

        ws.auto_filter.ref = ws.dimensions
        wb.save(filepath)
        return True

    def get_entry_count(self):
        return len(self.entries)

    def clear_all(self):
        self.entries.clear()
        self.current_entry.clear()
